import os
import re
from typing import List, Dict, Any
from pathlib import Path

def extract_text_from_pdf(file_path: str) -> str:
    import pypdf
    reader = pypdf.PdfReader(file_path)
    text_parts = []
    for i, page in enumerate(reader.pages):
        page_text = page.extract_text() or ""
        text_parts.append(f"\n[Page {i + 1}]\n{page_text}")
    return "\n".join(text_parts)

def extract_text_from_docx(file_path: str) -> str:
    import docx
    doc = docx.Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text:
                paragraphs.append(row_text)
    return "\n\n".join(paragraphs)

def extract_text_from_pptx(file_path: str) -> str:
    from pptx import Presentation
    prs = Presentation(file_path)
    slides_text = []
    for i, slide in enumerate(prs.slides):
        slide_parts = [f"[Slide {i + 1}]"]
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_parts.append(shape.text.strip())
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                slide_parts.append(f"(Notes: {notes})")
        slides_text.append("\n".join(slide_parts))
    return "\n\n".join(slides_text)

def extract_text_from_excel(file_path: str) -> str:
    import pandas as pd
    excel_file = pd.ExcelFile(file_path)
    sheets_text = []
    for sheet_name in excel_file.sheet_names:
        df = pd.read_excel(excel_file, sheet_name=sheet_name)
        sheets_text.append(f"[Sheet: {sheet_name}]\n" + df.to_markdown(index=False))
    return "\n\n".join(sheets_text)

def extract_text_from_file(file_path: str, gemini_api_key: str = None) -> str:
    path = Path(file_path)
    ext = path.suffix.lower()

    if ext in [".txt", ".md", ".markdown"]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    elif ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext in [".docx", ".doc"]:
        return extract_text_from_docx(file_path)
    elif ext in [".pptx", ".ppt"]:
        return extract_text_from_pptx(file_path)
    elif ext in [".xlsx", ".xls", ".csv"]:
        if ext == ".csv":
            import pandas as pd
            df = pd.read_csv(file_path)
            return df.to_markdown(index=False)
        return extract_text_from_excel(file_path)
    elif ext in [".png", ".jpg", ".jpeg", ".webp"]:
        # Multimodal OCR with Gemini Vision
        if gemini_api_key:
            import google.generativeai as genai
            from PIL import Image
            genai.configure(api_key=gemini_api_key)
            model = genai.GenerativeModel("gemini-1.5-flash")
            img = Image.open(file_path)
            response = model.generate_content([
                "Transcribe all study text, formulas, diagrams, and structures from this study material completely and accurately in clean markdown.",
                img
            ])
            return response.text
        return f"[Image file: {path.name}]"
    elif ext in [".mp3", ".wav", ".m4a", ".mp4"]:
        # Multimodal Audio/Video lecture transcription
        if gemini_api_key:
            import google.generativeai as genai
            genai.configure(api_key=gemini_api_key)
            audio_file = genai.upload_file(path=file_path)
            model = genai.GenerativeModel("gemini-1.5-flash")
            response = model.generate_content([
                "Transcribe and organize this lecture recording into detailed, structured study notes with timestamps, key concepts, and summaries.",
                audio_file
            ])
            return response.text
        return f"[Audio/Video lecture: {path.name}]"
    else:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

def chunk_text(text: str, filename: str, chunk_size: int = 1000, chunk_overlap: int = 200) -> List[Dict[str, Any]]:
    """
    Splits document into logical study chunks respecting markdown headers and chapters.
    """
    # Try splitting by chapter or major headers first
    header_splits = re.split(r'(?=\n#{1,3}\s+)', text)
    chunks = []
    chunk_index = 0

    current_topic = Path(filename).stem.replace("_", " ").title()

    for section in header_splits:
        section = section.strip()
        if not section:
            continue

        # Detect topic or chapter name from header
        match = re.match(r'^#{1,3}\s+(.+)', section)
        if match:
            current_topic = match.group(1).strip()

        # If section is small enough, keep as single chunk
        if len(section) <= chunk_size:
            chunk_index += 1
            chunks.append({
                "chunk_id": f"{filename}_{chunk_index}",
                "filename": filename,
                "topic": current_topic,
                "text": section
            })
        else:
            # Sub-split into overlapping windows
            start = 0
            while start < len(section):
                end = min(start + chunk_size, len(section))
                # Try breaking at newline or sentence
                if end < len(section):
                    last_break = max(section.rfind("\n", start, end), section.rfind(". ", start, end))
                    if last_break > start + 300:
                        end = last_break + 1

                sub_text = section[start:end].strip()
                if sub_text:
                    chunk_index += 1
                    chunks.append({
                        "chunk_id": f"{filename}_{chunk_index}",
                        "filename": filename,
                        "topic": current_topic,
                        "text": sub_text
                    })
                start = end - chunk_overlap if end < len(section) else end

    return chunks
